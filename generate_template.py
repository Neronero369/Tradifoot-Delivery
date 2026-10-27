#!/usr/bin/env python3
"""
Tradifoot Logistics - Master Delivery Program Template Generator
Creates a standardized, machine-readable Excel template: template_programme_de_livraison.xlsx
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def generate_master_template(output_path="template_programme_de_livraison.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Programme_Livraison"
    ws.views.sheetView[0].showGridLines = True

    # 1. Standard Column Headers (Row 1 - Strict Schema, No Merged Cells)
    headers = [
        "ID",
        "Date",
        "Magasin",
        "Chauffeur",
        "Matricule",
        "Quantite",
        "Status",
        "Remarques"
    ]
    ws.append(headers)

    # 2. Tradifoot Navy Corporate Palette & Styles
    navy_header_fill = PatternFill(start_color="0B1F3F", end_color="0B1F3F", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="0B1F3F"))

    # Style Header Row
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border

    # 3. Realistic Sample Rows
    sample_data = [
        [
            1001,
            "2026-10-28",
            "SKECHERS BAB EZZOUAR",
            "SAADI BILLEL",
            "01423-116-16",
            180,
            "Non Livré",
            "Livraison quai B avant 11h - Réceptionnaire: Ahmed"
        ],
        [
            1002,
            "2026-10-28",
            "GEOX ALGER CENTRE (DIDOUCHE)",
            "SNOUCI SEIF-EDDINE",
            "08912-118-16",
            240,
            "Non Livré",
            "Appeler le gérant 30 min avant arrivée"
        ],
        [
            1003,
            "2026-10-28",
            "CLARKS CC ORAN SENIA",
            "BOUKHALFA ADEL",
            "14205-121-31",
            320,
            "Non Livré",
            "Déchargement sous-sol réserve 2 - Accès véhicules autorisés"
        ]
    ]

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    normal_font = Font(name="Calibri", size=10, color="1E293B")
    id_font = Font(name="Consolas", size=10, bold=True, color="0B1F3F")
    qty_font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")

    for row_idx, row_values in enumerate(sample_data, start=2):
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 24
        is_even = (row_idx % 2 == 0)

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = cell_border
            cell.font = normal_font

            if is_even:
                cell.fill = zebra_fill

            # Specific column alignments & formatting
            if col_idx == 1:  # ID
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = id_font
            elif col_idx == 2:  # Date
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "YYYY-MM-DD"
            elif col_idx in (3, 8):  # Magasin, Remarques
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (4, 5):  # Chauffeur, Matricule
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 6:  # Quantité
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
                cell.font = qty_font
            elif col_idx == 7:  # Status
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 4. Data Validations
    # A. Authorized Chauffeurs Dropdown
    chauffeurs = [
        "SAADI BILLEL",
        "SNOUCI SEIF-EDDINE",
        "BOUKHALFA ADEL",
        "BALLA YOUCEF",
        "BAHLOULI ADEL"
    ]
    dv_chauffeur = DataValidation(
        type="list",
        formula1=f'"{",".join(chauffeurs)}"',
        allow_blank=True
    )
    dv_chauffeur.error = "Chauffeur non reconnu. Veuillez choisir un chauffeur autorisé dans la liste déroulante."
    dv_chauffeur.errorTitle = "Chauffeur Invalide"
    dv_chauffeur.prompt = "Sélectionnez le chauffeur affecté à la tournée."
    dv_chauffeur.promptTitle = "Affectation Chauffeur"
    ws.add_data_validation(dv_chauffeur)
    dv_chauffeur.add("D2:D1000")

    # B. Delivery Status Dropdown
    dv_status = DataValidation(
        type="list",
        formula1='"Non Livré,Livré"',
        allow_blank=False
    )
    dv_status.error = "Le statut doit être strictement 'Non Livré' ou 'Livré'."
    dv_status.errorTitle = "Statut Invalide"
    dv_status.prompt = "Sélectionnez l'état d'avancement de la livraison."
    dv_status.promptTitle = "Statut Livraison"
    ws.add_data_validation(dv_status)
    dv_status.add("G2:G1000")

    # C. Quantity Integer Validation (>= 1)
    dv_qty = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1="1",
        allow_blank=False
    )
    dv_qty.error = "La quantité doit être un nombre entier positif supérieur ou égal à 1."
    dv_qty.errorTitle = "Quantité Invalide"
    ws.add_data_validation(dv_qty)
    dv_qty.add("F2:F1000")

    # 5. Column Widths
    fixed_widths = {
        "A": 12,  # ID
        "B": 15,  # Date
        "C": 36,  # Magasin
        "D": 26,  # Chauffeur
        "E": 20,  # Matricule
        "F": 15,  # Quantité
        "G": 16,  # Status
        "H": 45   # Remarques
    }
    for col_letter, width in fixed_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Save Master Template
    wb.save(output_path)
    print(f"[OK] Master Delivery Template generated successfully: {output_path}")

if __name__ == "__main__":
    generate_master_template()
