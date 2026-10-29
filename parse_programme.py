#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tradifoot Logistics - Program Ingestion Parser, Validator & Upsert Engine
Production Excel Delivery Schedule Parser & Supabase Synchronization.
Dynamically locates headers (Boutique, Globale, Chauffeur, Livraison)
and maps records directly to the Supabase database schema.
"""

import sys
import os
import re
import json
import urllib.request
import urllib.error
from datetime import datetime, date
import openpyxl

DEFAULT_SUPABASE_REST_URL = "https://gmdgnzokgqyivlfeosdy.supabase.co/rest/v1/deliveries"
DEFAULT_SUPABASE_API_KEY = "sb_publishable_JPqzQ4b11U9BU8HzG16k2g_O2cHl5ur"

AUTHORIZED_CHAUFFEURS = {
    "SAADI BILLEL",
    "SNOUCI SEIF-EDDINE",
    "BOUKHALFA ADEL",
    "BALLA YOUCEF",
    "BAHLOULI ADEL"
}


def build_delivery_key(date_str, store_str, quantite):
    """Deterministic Composite Fingerprint Matching (Date + Magasin + Quantite)."""
    d = str(date_str or "").strip()
    m = " ".join(str(store_str or "").strip().upper().split())
    try:
        q = int(quantite)
    except (ValueError, TypeError):
        q = 0
    if not d or not m or q <= 0:
        return ""
    return f"{d}|{m}|{q}"


def normalize_composite_key(date_str, store_str, driver_str):
    """Generates a canonical composite key [Date + Magasin + Chauffeur]."""
    d = str(date_str or "").strip()
    m = " ".join(str(store_str or "").strip().upper().split())
    c = " ".join(str(driver_str or "").strip().upper().split())
    return f"{d}|||{m}|||{c}"


def normalize_date_to_iso(val, fallback=""):
    """
    Converts date representations (datetime, date, DD/MM/YYYY, etc.) to YYYY-MM-DD.
    """
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val or "").strip()
    if not s:
        return fallback or datetime.today().strftime("%Y-%m-%d")

    # Priority 1: DD/MM/YYYY and variations
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    # Priority 2: ISO formatted timestamps
    try:
        return datetime.fromisoformat(s[:19]).strftime("%Y-%m-%d")
    except Exception:
        pass

    return s[:10]


def parse_delivery_excel(file_path):
    """
    Parses a production Programme de Livraison Excel workbook dynamically by header name.

    CRITICAL COLUMN MAPPING:
      - Boutique   -> maps to magasin (String, uppercase, trimmed)
      - Globale    -> maps to quantite (Integer, default 0, skip row if 0/empty)
      - Chauffeur  -> maps to chauffeur (String, uppercase, default 'NON SPÉCIFIÉ')
      - Livraison  -> maps to date (Convert from DD/MM/YYYY to YYYY-MM-DD)

    Completely ignores all operational columns (Wilaya, Region, ACCESSOIRES,
    CHAUSSETTES, SAC A DOS, FOOTWEAR, Sachet, P, OM, Retour, Chargement,
    Materiel, REMARQUE).

    Generates composite delivery_id (LIV-YYYYMMDD-STORE-CHAUF-INDEX).
    Sets status strictly to 'Non Livré' for all newly parsed rows.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # Locate the delivery program sheet (or default to active sheet)
    ws = wb.active
    for s_name in wb.sheetnames:
        if re.search(r'prog|poro|livr', s_name, re.I):
            ws = wb[s_name]
            break

    # 1. Dynamic Header Resolution (search rows 1 to 7)
    header_row_idx = None
    col_boutique = None
    col_globale = None
    col_chauffeur = None
    col_livraison = None
    raw_headers = []

    for r_idx, row in enumerate(ws.iter_rows(max_row=7, values_only=True), start=1):
        clean_row = [str(c or "").strip().lower() for c in row]
        has_boutique = any(h in ("boutique", "magasin") or "boutique" in h or "magasin" in h for h in clean_row)
        has_globale = any(h in ("globale", "global", "quantite", "quantité", "qte") or "globale" in h or "quantite" in h for h in clean_row)
        
        if has_boutique and (has_globale or any("livraison" in h or "chauffeur" in h for h in clean_row)):
            header_row_idx = r_idx
            raw_headers = [str(c or "").strip() for c in row]
            for c_idx, h in enumerate(clean_row):
                if h in ("boutique", "magasin") or ("boutique" in h and col_boutique is None) or ("magasin" in h and col_boutique is None):
                    col_boutique = c_idx
                elif h in ("globale", "global", "quantite", "quantité", "qte") or ("globale" in h and col_globale is None):
                    col_globale = c_idx
                elif h in ("chauffeur", "driver") or ("chauffeur" in h and col_chauffeur is None):
                    col_chauffeur = c_idx
                elif h in ("livraison", "date") or ("livraison" in h and col_livraison is None):
                    col_livraison = c_idx
            break

    if header_row_idx is None or col_boutique is None:
        wb.close()
        raise ValueError(
            f"Could not locate required 'Boutique' column in {file_path}. Headers detected: {raw_headers}"
        )

    records = []
    errors = []
    record_idx = 1

    # 2. Iterate data rows starting immediately below the detected header row
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
        if not any(row):
            continue

        # Column 1: Boutique -> magasin (String, uppercase, trimmed)
        raw_store = str(row[col_boutique] or "").strip() if col_boutique < len(row) and row[col_boutique] is not None else ""
        magasin = " ".join(raw_store.upper().split())

        # Column 2: Globale -> quantite (Integer, default 0 if blank)
        raw_qty = row[col_globale] if col_globale is not None and col_globale < len(row) else 0
        try:
            if raw_qty is not None and str(raw_qty).strip() != "":
                qty = int(float(str(raw_qty).replace(",", ".")))
            else:
                qty = 0
        except (ValueError, TypeError):
            qty = 0

        # Skip rule: If a row has an empty Boutique or if Globale is 0/empty, skip entirely
        if not magasin or magasin.startswith("TOTAL") or magasin == "NONE" or qty <= 0:
            continue

        # Column 3: Chauffeur -> chauffeur (String, uppercase, default 'NON SPÉCIFIÉ' if blank)
        raw_driver = str(row[col_chauffeur] or "").strip() if col_chauffeur is not None and col_chauffeur < len(row) and row[col_chauffeur] is not None else ""
        chauffeur = " ".join(raw_driver.upper().split()) if raw_driver else "NON SPÉCIFIÉ"
        if chauffeur in ("", "NONE", "NULL", "-", "/"):
            chauffeur = "NON SPÉCIFIÉ"

        # Column 4: Livraison -> date (Convert from DD/MM/YYYY to YYYY-MM-DD)
        raw_date = row[col_livraison] if col_livraison is not None and col_livraison < len(row) else None
        normalized_date = normalize_date_to_iso(raw_date)

        # Composite delivery_id: LIV-YYYYMMDD-STORE-CHAUF-INDEX
        ymd = normalized_date.replace("-", "")
        store_slug = re.sub(r'[^A-Z0-9]', '', magasin)[:12]
        chauf_slug = re.sub(r'[^A-Z0-9]', '', chauffeur)[:8]
        delivery_id = f"LIV-{ymd}-{store_slug}-{chauf_slug}-{record_idx:04d}"

        # Matched schema payload for Supabase
        record = {
            "id": delivery_id,
            "delivery_id": delivery_id,
            "date": normalized_date,
            "magasin": magasin,
            "chauffeur": chauffeur,
            "matricule": "",
            "quantite": qty,
            "status": "Non Livré",
            "delivery_key": build_delivery_key(normalized_date, magasin, qty)
        }
        records.append(record)
        record_idx += 1

    wb.close()
    return {
        "success": len(errors) == 0,
        "total_records": len(records),
        "records": records,
        "errors": errors
    }


def upsert_manifest_records(incoming_records, existing_records):
    """
    Performs in-memory Upsert logic with strict Status Preservation.

    CRITICAL STATUS PRESERVATION CONTRACT:
    - If a delivery already exists and its current status is 'Livré', PRESERVE 'Livré'.
      Never downgrade to 'Non Livré' or empty values.
    - If a delivery exists with status 'Non Livré', update mutable fields (Quantité, Matricule)
      but keep status as 'Non Livré' (unless incoming explicitly provides 'Livré').
    - If a delivery is brand new, insert as 'Non Livré'.
    """
    id_map = {}
    comp_map = {}

    for rec in existing_records:
        r_id = str(rec.get("id") or "").strip().upper()
        if r_id:
            id_map[r_id] = rec

        ckey = normalize_composite_key(rec.get("date"), rec.get("magasin"), rec.get("chauffeur"))
        if ckey:
            comp_map[ckey] = rec

    stats = {
        "total": len(incoming_records),
        "preserved_livre": 0,
        "upgraded_livre": 0,
        "updated_non_livre": 0,
        "inserted_new": 0
    }

    merged = list(existing_records)

    for item in incoming_records:
        in_id = str(item.get("id") or "").strip().upper()
        ckey = normalize_composite_key(item.get("date"), item.get("magasin"), item.get("chauffeur"))

        matched = id_map.get(in_id) or comp_map.get(ckey)

        if matched:
            curr_status = matched.get("status", "Non Livré")
            incoming_status = item.get("status", "Non Livré")

            if curr_status == "Livré":
                final_status = "Livré"
                stats["preserved_livre"] += 1
            elif incoming_status == "Livré":
                final_status = "Livré"
                stats["upgraded_livre"] += 1
            else:
                final_status = "Non Livré"
                stats["updated_non_livre"] += 1

            matched["status"] = final_status
            matched["completed"] = (final_status == "Livré")
            if item.get("quantite"):
                matched["quantite"] = item["quantite"]
            if item.get("matricule"):
                matched["matricule"] = item["matricule"]
        else:
            new_item = dict(item)
            if new_item.get("status") != "Livré":
                new_item["status"] = "Non Livré"
            new_item["completed"] = (new_item["status"] == "Livré")
            merged.append(new_item)
            stats["inserted_new"] += 1

            if in_id:
                id_map[in_id] = new_item
            if ckey:
                comp_map[ckey] = new_item

    return merged, stats


def fetch_supabase_records(supabase_url=None, api_key=None):
    """
    Fetches all current deliveries from Supabase REST endpoint to protect existing
    completed ('Livré') statuses and avoid duplicate processing.
    """
    url = (supabase_url or DEFAULT_SUPABASE_REST_URL).split("?")[0]
    key = api_key or DEFAULT_SUPABASE_API_KEY
    endpoint = f"{url}?select=delivery_id,delivery_key,date,magasin,chauffeur,quantite,status,delivered_by"
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}"
    }
    try:
        req = urllib.request.Request(endpoint, headers=headers)
        with urllib.request.urlopen(req, timeout=30) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except Exception as e:
        print(f"[!] Note: fetch_supabase_records encountered notice ({e})")
        return []


def push_manifest_to_supabase(records, supabase_url=None, api_key=None, only_new=True):
    """
    Pushes delivery records to Supabase REST endpoint via POST upsert
    with STRICT STATUS PRESERVATION & DEDUPLICATION:

    1. Queries existing records in Supabase before pushing.
    2. Any existing delivery that is already 'Livré' is NEVER overwritten with 'Non Livré'.
    3. If only_new is True, skips records that already exist in Supabase so old delivery
       programs never get re-processed or override database records.
    4. Only truly new deliveries (or non-duplicate updates) are pushed.
    """
    url = supabase_url or DEFAULT_SUPABASE_REST_URL
    key = api_key or DEFAULT_SUPABASE_API_KEY

    if isinstance(records, str):
        parsed = parse_delivery_excel(records)
        records = parsed.get("records", [])

    if not records:
        print("[!] No records provided to push to Supabase.")
        return {"status": "success", "count": 0, "data": []}

    # 1. Fetch live Supabase state for status preservation
    existing_records = fetch_supabase_records(url, key)
    existing_by_key = {}
    existing_by_id = {}
    existing_by_comp = {}

    for row in existing_records:
        dkey = str(row.get("delivery_key") or "").strip()
        did = str(row.get("delivery_id") or row.get("id") or "").strip().upper()
        r_date = str(row.get("date") or "").strip()
        r_mag = " ".join(str(row.get("magasin") or "").strip().upper().split())
        r_drv = " ".join(str(row.get("chauffeur") or "").strip().upper().split())

        if dkey:
            existing_by_key[dkey] = row
        if did:
            existing_by_id[did] = row
        if r_date and r_mag:
            existing_by_comp[(r_date, r_mag, r_drv)] = row
            existing_by_comp[(r_date, r_mag)] = row

    payload = []
    skipped_existing = 0
    preserved_livre = 0

    for r in records:
        rec_id = r.get("id") or r.get("delivery_id")
        if not rec_id:
            continue
        try:
            qty = int(r.get("quantite", 0))
        except (ValueError, TypeError):
            qty = 0

        deliv_date = str(r.get("date") or "").strip()
        store = str(r.get("magasin") or "").strip()
        driver = str(r.get("chauffeur") or "").strip()
        matricule = str(r.get("matricule") or "").strip()
        raw_status = str(r.get("status") or "Non Livré").strip()

        dkey = r.get("delivery_key") or build_delivery_key(deliv_date, store, qty)

        norm_comp = (deliv_date, " ".join(store.upper().split()), " ".join(driver.upper().split()))
        norm_comp_store = (deliv_date, " ".join(store.upper().split()))

        existing_row = (
            existing_by_key.get(dkey)
            or existing_by_id.get(str(rec_id).strip().upper())
            or existing_by_comp.get(norm_comp)
            or existing_by_comp.get(norm_comp_store)
        )

        if existing_row:
            s_stat = str(existing_row.get("status") or "")
            is_already_livre = bool("livr" in s_stat.lower() and "non" not in s_stat.lower())

            if is_already_livre:
                preserved_livre += 1
                # CRITICAL: Completed delivery in Supabase MUST NOT be downgraded to 'Non Livré'
                if only_new:
                    skipped_existing += 1
                    continue
                final_status = "Livré"
            else:
                if only_new:
                    # Already present in Supabase, skip to avoid overrun
                    skipped_existing += 1
                    continue
                final_status = raw_status
        else:
            final_status = raw_status

        mapped = {
            "delivery_id": str(rec_id).strip(),
            "date": deliv_date,
            "magasin": store,
            "chauffeur": driver,
            "matricule": matricule,
            "quantite": qty,
            "status": final_status,
            "delivery_key": dkey
        }
        payload.append(mapped)

    if skipped_existing > 0:
        print(f"[i] Status Preservation: Skipped {skipped_existing} existing deliveries ({preserved_livre} preserved as 'Livré').")

    if not payload:
        print("[i] All incoming records already exist in Supabase. Nothing new to push.")
        return {"status": "success", "count": 0, "message": "All records up to date"}

    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"
    }

    # Ensure on_conflict target matches the table unique constraint (delivery_key)
    if "on_conflict=" not in url:
        sep = "&" if "?" in url else "?"
        endpoint_url = f"{url}{sep}on_conflict=delivery_key"
    else:
        endpoint_url = url

    BATCH_SIZE = 100
    total_pushed = 0
    try:
        for i in range(0, len(payload), BATCH_SIZE):
            batch = payload[i:i + BATCH_SIZE]
            req_data = json.dumps(batch).encode("utf-8")
            req = urllib.request.Request(endpoint_url, data=req_data, headers=headers, method="POST")
            try:
                with urllib.request.urlopen(req, timeout=45) as resp:
                    pass
            except urllib.error.HTTPError as he:
                # Fallback to direct url if endpoint rejects on_conflict parameter
                if he.code == 400 and endpoint_url != url:
                    fallback_req = urllib.request.Request(url, data=req_data, headers=headers, method="POST")
                    with urllib.request.urlopen(fallback_req, timeout=45) as resp:
                        pass
                else:
                    raise he
            total_pushed += len(batch)

        print(f"[+] Successfully pushed {total_pushed} new record(s) to Supabase.")
        return {"status": "success", "count": total_pushed}
    except urllib.error.HTTPError as e:
        err_body = e.read().decode("utf-8", errors="ignore")
        print(f"[-] Supabase HTTP Error {e.code}: {err_body}")
        return {"status": "error", "code": e.code, "message": str(e), "details": err_body}
    except urllib.error.URLError as e:
        print(f"[-] Supabase URL Error: {e}")
        return {"status": "error", "message": str(e)}
    except Exception as e:
        print(f"[-] Supabase Push Unexpected Error: {e}")
        return {"status": "error", "message": str(e)}


if __name__ == "__main__":
    test_files = [
        "data/programs/20261029_1423_PROGRAMME 29-10-2026.xlsx",
        "template_programme_de_livraison.xlsx"
    ]
    for tf in test_files:
        if os.path.exists(tf):
            print(f"\n--- Testing parse_delivery_excel on: {tf} ---")
            res = parse_delivery_excel(tf)
            print(f"Success: {res['success']}")
            print(f"Total Rows: {res['total_records']}")
            print(f"Sample Parsed Record: {res['records'][0] if res['records'] else None}")
